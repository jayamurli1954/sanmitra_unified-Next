#!/usr/bin/env python3
"""One-off: add 'HR & Payroll' and 'Manufacturing & Cost Centres' sections to
docs/MitraBooks_User_Manual.docx (and emit a markdown copy for the repo).

Inserts the two enterprise-module sections before '22. Workflow Tips' and renumbers
the four tail sections (22-25 -> 24-27). Reuses the manual's own styles and list
numbering (numId 3 = numbered steps, numId 4 = bullets) so the new sections look
identical to the rest of the manual. The Table of Contents is an auto-field — it
refreshes in Word (Ctrl+A then F9, or on open).
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
MANUAL = ROOT / "docs" / "MitraBooks_User_Manual.docx"
ASSET_COPY = ROOT / "frontend" / "assets" / "mitrabooks" / "MitraBooks_User_Manual.docx"
MD_OUT = ROOT / "docs" / "MITRABOOKS_HR_AND_MANUFACTURING_USER_GUIDE.md"

STEP_NUMID = 3     # decimal numbered steps
BULLET_NUMID = 4   # bullets

# Content model: ("h1"|"h2"|"p"|"step"|"bullet", runs)
# runs = list of (text, bold). Plain string is shorthand for [(text, False)].
def R(*parts):
    out = []
    for p in parts:
        out.append(p if isinstance(p, tuple) else (p, False))
    return out


CONTENT = [
    # ===================== HR & PAYROLL =====================
    ("h1", "22. HR & Payroll"),
    ("p", R("HR & Payroll is an ", ("enterprise add-on", True),
            ". It is off by default and must be provisioned by your platform "
            "administrator and then enabled in MitraBooks before the workspace becomes "
            "active. It manages employees, salary structures, payroll runs, leave, "
            "investment declarations (Form 12BB) and full & final settlements, and posts "
            "payroll to the General Ledger.")),
    ("p", R("Open it from ", ("Human Resources → HR & Payroll", True),
            " in the left navigation.")),

    ("h2", "22.1 Activating HR & Payroll"),
    ("step", R("Your platform administrator first ", ("provisions", True),
               " the add-on for your organization.")),
    ("step", R("A business administrator then sees an ", ("Enable HR & Payroll", True),
               " button in the workspace and clicks it.")),
    ("step", R("Once active, the tabs Employees, Payroll, Leave, Form 12BB, Full & Final "
               "and Analytics appear.")),
    ("p", R("Access is role-based: an HR Manager or the tenant administrator can manage; "
            "a Payroll Auditor has read-only access.")),

    ("h2", "22.2 Salary Structures"),
    ("p", R("A salary structure is a reusable template of formula-driven components "
            "(Basic, HRA, allowances) used when assigning salaries.")),
    ("step", R("On the ", ("Employees", True), " tab, find ", ("Salary Structures", True), ".")),
    ("step", R("Enter a name (e.g. ", ("Standard (Non-metro)", True),
               "), a Basic formula (e.g. ", ("GROSS * 0.5", True),
               ") and an HRA formula (e.g. ", ("BASIC * 0.4", True), ").")),
    ("step", R("Click ", ("Add Structure", True), ".")),

    ("h2", "22.3 Adding an Employee"),
    ("p", R("Employees follow an onboarding lifecycle: ", ("Offered", True),
            " → ", ("Joined", True), " or ", ("Declined", True),
            ". An employee code is minted only when the candidate joins, so declined "
            "candidates never consume a code.")),
    ("step", R("Click ", ("+ Add Employee", True),
               " and fill in name, designation, dates and contact details. The new "
               "employee starts in the Offered state.")),
    ("step", R("Use ", ("Assign Salary", True),
               " to attach a salary structure and the CTC figures.")),
    ("step", R("When the candidate joins, click ", ("Mark Joined", True),
               " — the employee code is generated at this point.")),

    ("h2", "22.4 Appointment & Joining Letters"),
    ("step", R("Open ", ("Letter Settings", True),
               " to configure the clauses and branding that appear on letters.")),
    ("step", R("From an employee row, generate the ", ("Appointment Letter", True),
               " (offer stage) or the ", ("Joining Letter", True),
               " (after joining) as a PDF.")),

    ("h2", "22.5 Running Payroll"),
    ("step", R("Go to the ", ("Payroll", True), " tab and click ", ("Run Payroll", True),
               " for the pay period.")),
    ("step", R("The engine computes each slip with EPF, ESI, Professional Tax (state "
               "slabs), TDS (new and old regime, with rebate and cess) and gratuity, all "
               "in fixed-precision money.")),
    ("step", R("Loss-of-pay is derived from approved leave — it is never typed in.")),
    ("step", R("The run produces salary slips and one consolidated journal entry to the "
               "GL. Download any slip as a PDF from the run.")),

    ("h2", "22.6 Leave Management"),
    ("step", R("Create leave types on the ", ("Leave", True),
               " tab (mark a type as loss-of-pay if applicable).")),
    ("step", R("Allocate leave balances to employees.")),
    ("step", R("Employees apply for leave; an approver approves or rejects. Approved "
               "leave feeds the loss-of-pay calculation in payroll.")),

    ("h2", "22.7 Form 12BB (Investment Declarations)"),
    ("p", R("Form 12BB lets employees declare investments and submit proofs so old-regime "
            "TDS is computed correctly.")),
    ("step", R("On the ", ("Form 12BB", True),
               " tab the employee declares investments and uploads proof.")),
    ("step", R("HR verifies the declaration; verified amounts flow into the old-regime "
               "TDS calculation at payroll time.")),

    ("h2", "22.8 Full & Final Settlement"),
    ("step", R("On the ", ("Full & Final", True),
               " tab, create an F&F for a leaving employee. Gratuity is computed from "
               "tenure.")),
    ("step", R("Move it through ", ("Draft → Approved → Paid", True),
               ". Download the F&F statement as a PDF.")),

    ("h2", "22.9 Analytics"),
    ("p", R("The ", ("Analytics", True),
            " tab shows a trailing-month payroll dashboard (headcount, payroll cost and "
            "statutory totals) computed from the posted runs.")),

    # ============== MANUFACTURING & COST CENTRES ==============
    ("h1", "23. Manufacturing & Cost Centres"),
    ("p", R("This is an ", ("enterprise add-on", True),
            " with two independent layers: ", ("Cost-Centre Accounting", True),
            " (departmental/branch budgets and P&L) and ", ("Manufacturing", True),
            " (bills of materials and work orders). Manufacturing depends on cost "
            "centres. Both are off by default and provisioned per organization.")),
    ("p", R("Open it from ", ("Manufacturing → Manufacturing", True),
            " in the left navigation. It has five tabs: Cost Centres, Budgets, "
            "Cost-Centre P&L, BOMs and Work Orders.")),
    ("p", R(("Important — how it affects your books: ", True),
            "the module is built for periodic inventory. Work orders do not post their "
            "own inventory journals (that would double-count against the period-end "
            "closing-stock entry). Instead they record production and a standard-vs-actual "
            "variance for reporting, and feed finished goods and raw-material consumption "
            "into the stock register so closing-stock valuation stays correct.")),

    ("h2", "23.1 Activating the Add-on"),
    ("step", R("The platform administrator provisions Cost-Centre Accounting (and, if "
               "needed, Manufacturing) for the organization.")),
    ("step", R("A business administrator clicks ", ("Enable Cost Centres", True),
               "; enabling ", ("Manufacturing", True),
               " automatically requires cost centres to be on.")),

    ("h2", "23.2 Cost Centres"),
    ("p", R("A cost centre is a department, branch or activity you want to measure "
            "separately. Cost centres can be nested (e.g. Assembly Line rolls up into "
            "Factory, which rolls up into Operations).")),
    ("step", R("On the ", ("Cost Centres", True),
               " tab, enter a Code (e.g. MFG-ASY-01), a Name (e.g. Assembly Line 1) and "
               "an optional Parent code.")),
    ("step", R("Click ", ("+ Add Cost Centre", True),
               ". The hierarchy is shown beneath the list.")),

    ("h2", "23.3 Tagging Postings to a Cost Centre"),
    ("p", R("Cost centres become useful when transactions carry them. A journal line may "
            "be tagged with a cost centre; the system rejects any cost centre that does "
            "not belong to your entity, so figures never mix across tenants or entities.")),

    ("h2", "23.4 Cost-Centre Budgets"),
    ("step", R("On the ", ("Budgets", True),
               " tab, pick a cost centre, a fiscal year (and optional month), an account "
               "and an allocated amount, then ", ("+ Add Budget", True), ".")),
    ("step", R("Move a budget ", ("Draft → Approved → Locked", True),
               ". Only approved/locked budgets drive the variance report.")),
    ("step", R("Click ", ("Vs Actual", True),
               " on a budget to see allocated vs actual spend, variance and burn-rate per "
               "account, with unbudgeted spend surfaced separately.")),

    ("h2", "23.5 Cost-Centre P&L"),
    ("step", R("On the ", ("Cost-Centre P&L", True),
               " tab, choose a date range and click ", ("Run", True), ".")),
    ("step", R("The report shows income, expense and net per cost centre, plus an "
               "Untagged bucket, with totals that tie back to the period P&L.")),
    ("step", R("Export to ", ("CSV", True), " or ", ("Excel", True),
               " with the buttons provided.")),

    ("h2", "23.6 Bills of Materials (BOM)"),
    ("p", R("A BOM is the recipe for a finished good: the component items it consumes "
            "(with a standard rate and scrap allowance) and the operations performed "
            "(with an overhead rate). From these the system computes a deterministic "
            "standard cost. BOMs reference items from the inventory item master, so "
            "create your items first.")),
    ("step", R("On the ", ("BOMs", True), " tab, open ", ("+ New BOM", True),
               ", choose the finished good and an output quantity.")),
    ("step", R("Add each component (item, quantity, rate, scrap %) with ", ("Add line", True), ".")),
    ("step", R("Click ", ("Save BOM", True),
               ". The standard cost (total and per unit) is shown in the BOM list.")),

    ("h2", "23.7 Work Orders"),
    ("p", R("A work order plans production of a finished good against a BOM and tracks it "
            "through its lifecycle: ", ("Draft → Released → In Progress → Completed", True),
            " (or Cancelled).")),
    ("step", R("On the ", ("Work Orders", True),
               " tab, pick a BOM, a planned quantity and (optionally) the production cost "
               "centre, then ", ("+ Create Work Order", True),
               ". The standard cost is snapshotted.")),
    ("step", R("Use ", ("Release", True), " and ", ("Start", True),
               " to advance the work order.")),
    ("step", R("Click ", ("Complete", True),
               ", enter the produced quantity, actual overhead and actual material "
               "consumed (item, quantity, rate), then ", ("Confirm Completion", True), ".")),
    ("step", R("The work order shows the ", ("variance", True),
               " (actual vs standard); a favourable variance is marked. The finished "
               "goods and consumed materials flow into the stock register.")),

    ("h2", "23.8 How It Affects Stock and the Ledger"),
    ("bullet", R("Finished goods produced are added to stock at their production cost.")),
    ("bullet", R("Raw materials consumed reduce stock at weighted-average cost.")),
    ("bullet", R("No separate work-order inventory journal is posted; financial "
                 "recognition flows through the existing closing-stock entry, keeping the "
                 "books consistent under periodic inventory.")),
    ("bullet", R("Variance is management information for cost control, tagged to the "
                 "production cost centre.")),
]


def _set_numbering(p, num_id, ilvl=0):
    pPr = p._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    e_ilvl = OxmlElement("w:ilvl"); e_ilvl.set(qn("w:val"), str(ilvl))
    e_num = OxmlElement("w:numId"); e_num.set(qn("w:val"), str(num_id))
    numPr.append(e_ilvl); numPr.append(e_num)
    pPr.append(numPr)


def _runs(p, runs):
    if isinstance(runs, str):
        runs = [(runs, False)]
    for text, bold in runs:
        r = p.add_run(text)
        if bold:
            r.bold = True


def build_and_insert(doc) -> None:
    # Find the anchor heading "22. Workflow Tips" in the body (not the TOC cache).
    anchor = None
    for para in doc.paragraphs:
        if para.style and para.style.name == "Heading 1" and para.text.strip() == "22. Workflow Tips":
            anchor = para
            break
    if anchor is None:
        raise SystemExit("Could not find the '22. Workflow Tips' heading to insert before.")

    new_paras = []
    for kind, runs in CONTENT:
        if kind == "h1":
            p = doc.add_paragraph(style="Heading 1"); _runs(p, runs)
        elif kind == "h2":
            p = doc.add_paragraph(style="Heading 2"); _runs(p, runs)
        elif kind == "p":
            p = doc.add_paragraph(style="Normal"); _runs(p, runs)
        elif kind == "step":
            p = doc.add_paragraph(style="List Paragraph"); _runs(p, runs); _set_numbering(p, STEP_NUMID)
        elif kind == "bullet":
            p = doc.add_paragraph(style="List Paragraph"); _runs(p, runs); _set_numbering(p, BULLET_NUMID)
        else:
            continue
        new_paras.append(p)

    # add_paragraph appended them at the end of the body; move each before the anchor,
    # preserving order.
    for p in new_paras:
        anchor._p.addprevious(p._p)

    # Renumber the tail sections 22-25 -> 24-27 (body headings only; TOC auto-updates).
    renumber = {
        "22. Workflow Tips": "24. Workflow Tips",
        "23. Common Errors & Fixes": "25. Common Errors & Fixes",
        "24. Glossary": "26. Glossary",
        "25. Support": "27. Support",
    }
    for para in doc.paragraphs:
        if para.style and para.style.name == "Heading 1":
            t = para.text.strip()
            if t in renumber and para.runs:
                # Rewrite text in the first run, clear the rest.
                para.runs[0].text = renumber[t]
                for r in para.runs[1:]:
                    r.text = ""


def emit_markdown() -> None:
    lines = ["# MitraBooks — HR & Payroll and Manufacturing & Cost Centres (User Guide)",
             "",
             "*This is the source for the two enterprise-module sections added to "
             "`docs/MitraBooks_User_Manual.docx` (sections 22 and 23).*", ""]
    for kind, runs in CONTENT:
        if isinstance(runs, str):
            runs = [(runs, False)]
        text = "".join((f"**{t}**" if b else t) for t, b in runs)
        if kind == "h1":
            lines += ["", f"## {text}", ""]
        elif kind == "h2":
            lines += ["", f"### {text}", ""]
        elif kind == "p":
            lines += [text, ""]
        elif kind == "step":
            lines.append(f"1. {text}")
        elif kind == "bullet":
            lines.append(f"- {text}")
    MD_OUT.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"Wrote {MD_OUT.relative_to(ROOT)}")


def main() -> None:
    doc = Document(str(MANUAL))
    build_and_insert(doc)
    doc.save(str(MANUAL))
    print(f"Updated {MANUAL.relative_to(ROOT)}")
    if ASSET_COPY.parent.exists():
        import shutil
        shutil.copyfile(MANUAL, ASSET_COPY)
        print(f"Synced {ASSET_COPY.relative_to(ROOT)}")
    emit_markdown()


if __name__ == "__main__":
    main()
