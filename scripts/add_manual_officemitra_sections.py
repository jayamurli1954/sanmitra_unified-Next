#!/usr/bin/env python3
"""One-off: add 'OfficeMitra AI' section to docs/MitraBooks_User_Manual.docx.

Inserts section 24 before '24. Workflow Tips' and renumbers the four tail sections
(24-27 -> 25-28). Idempotent: skips if '24. OfficeMitra AI' already exists.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
MANUAL = ROOT / "docs" / "MitraBooks_User_Manual.docx"
ASSET_COPY = ROOT / "frontend" / "assets" / "mitrabooks" / "MitraBooks_User_Manual.docx"
MD_OUT = ROOT / "docs" / "MITRABOOKS_OFFICEMITRA_AI_USER_GUIDE.md"

STEP_NUMID = 3
BULLET_NUMID = 4


def R(*parts):
    out = []
    for p in parts:
        out.append(p if isinstance(p, tuple) else (p, False))
    return out


CONTENT = [
    ("h1", "24. OfficeMitra AI"),
    (
        "p",
        R(
            ("OfficeMitra AI", True),
            " is SanMitra's thin AI productivity layer inside MitraBooks ERP. It helps you "
            "capture tasks, summarize pasted email and meeting notes, parse calendar text, "
            "view in-app notifications, and generate a ",
            ("Today Brief", True),
            ". When enabled for your tenant, it can also run the ",
            ("CA Analysis Pack (MIS)", True),
            " workflow: import Excel facts, reconcile with maker/checker, and export "
            "PDF/Excel/PPT packs. It does ",
            ("not", True),
            " post journals, invoices, or GST entries from MIS flows.",
        ),
    ),
    (
        "p",
        R(
            "Open it from ",
            ("Main Workspaces → OfficeMitra AI", True),
            " in the left navigation, or from the ",
            ("OfficeMitra AI", True),
            " chip on the dashboard.",
        ),
    ),
    (
        "p",
        R(
            "All AI output is ",
            ("advisory only", True),
            " — review before acting. It is not final legal or financial advice.",
        ),
    ),

    ("h2", "24.1 Opening the workspace"),
    (
        "step",
        R(
            "Sign in to MitraBooks ERP with a user whose tenant has the ",
            ("office_ai", True),
            " module enabled.",
        ),
    ),
    (
        "step",
        R("In the sidebar, under ", ("Main Workspaces", True), ", click ", ("OfficeMitra AI", True), "."),
    ),
    (
        "step",
        R(
            "The workspace opens with tabs across the top: ",
            ("Tasks", True),
            ", ",
            ("Email Summary", True),
            ", ",
            ("Calendar", True),
            ", ",
            ("Meeting Notes", True),
            ", ",
            ("Notifications", True),
            ", and ",
            ("Today Brief", True),
            ". Additional tabs appear when your administrator enables write-back, MIS, or workflows.",
        ),
    ),

    ("h2", "24.2 Tasks"),
    (
        "p",
        R(
            "Use Tasks to track follow-ups manually or with AI assistance. Paste notes or "
            "instructions in the text area, then:",
        ),
    ),
    ("step", R("Click ", ("Save as task", True), " to store the first line as a manual task.")),
    (
        "step",
        R(
            "Click ",
            ("Generate with AI", True),
            " to ask OfficeMitra to suggest tasks (requires AI provider configuration; "
            "otherwise the UI soft-fails with a clear message).",
        ),
    ),
    ("step", R("Click ", ("Done", True), " on an open task when it is complete.")),
    ("step", R("Click ", ("Refresh", True), " to reload the task list.")),

    ("h2", "24.3 Email Summary"),
    (
        "p",
        R(
            "Paste the full text of an email (no mailbox OAuth in this release). Click ",
            ("Summarize + suggest tasks", True),
            " to store a summary and optional follow-up tasks.",
        ),
    ),

    ("h2", "24.4 Calendar (paste-in)"),
    (
        "p",
        R(
            "Paste ICS calendar blocks or simple agenda lines (for example ",
            ('"10:00 Client GST review"', True),
            "). Click ",
            ("Parse + save events", True),
            " to add today's events. Use ",
            ("Refresh today", True),
            " to reload the list.",
        ),
    ),

    ("h2", "24.5 Meeting Notes"),
    (
        "p",
        R(
            "Paste meeting notes, then click ",
            ("Summarize + suggest tasks", True),
            ". Summaries are stored for later reference; optional tasks may be proposed when write-back is enabled.",
        ),
    ),

    ("h2", "24.6 Notifications"),
    (
        "p",
        R(
            "OfficeMitra maintains an in-app notification inbox (not email push). Click ",
            ("Mark read", True),
            " on individual items or ",
            ("Refresh", True),
            " to update the unread count badge on the tab.",
        ),
    ),

    ("h2", "24.7 Today Brief"),
    (
        "p",
        R(
            "The Today Brief combines OfficeMitra-native data and any available connector "
            "sections (for example open tasks and today's calendar). Click ",
            ("Generate today's brief", True),
            " to create or refresh today's digest. Click ",
            ("Load latest", True),
            " to fetch the saved brief without regenerating.",
        ),
    ),

    ("h2", "24.8 Proposals (confirmation and checker)"),
    (
        "p",
        R(
            "When ",
            ("office_ai.writeback", True),
            " or MIS export/reconcile is enabled, AI-suggested writes and MIS actions appear "
            "in the ",
            ("Proposals", True),
            " tab instead of executing immediately.",
        ),
    ),
    ("step", R("Review the summary, action type, confidence, and status.")),
    (
        "step",
        R(
            "As ",
            ("maker", True),
            ", click ",
            ("Confirm", True),
            " on pending proposals you prepared.",
        ),
    ),
    (
        "step",
        R(
            "When maker-checker policy applies, a different user (",
            ("checker", True),
            ") clicks ",
            ("Approve (checker)", True),
            " — you cannot approve your own proposal.",
        ),
    ),
    ("step", R("Click ", ("Dismiss", True), " to discard a proposal you do not want to execute.")),

    ("h2", "24.9 MIS Packs (CA Analysis Pack)"),
    (
        "p",
        R(
            "Available when ",
            ("office_ai.mis", True),
            " and related import/export flags are enabled. This is for monthly MIS assembly — "
            "not day-to-day bookkeeping.",
        ),
    ),
    ("step", R("Open the ", ("MIS Packs", True), " tab.")),
    (
        "step",
        R(
            "Choose a metric pack and period, then click ",
            ("Create draft pack", True),
            ".",
        ),
    ),
    (
        "step",
        R(
            "Upload the SanMitra CA MIS Excel template. Use ",
            ("Persist valid rows after validation", True),
            " when ready to insert facts.",
        ),
    ),
    (
        "step",
        R(
            "Review the KPI dashboard and fact table derived from imported rows (not AI estimates).",
        ),
    ),
    (
        "step",
        R(
            "As maker, enter an optional data-quality score and click ",
            ("Submit reconcile", True),
            ".",
        ),
    ),
    (
        "step",
        R(
            "After checker approval, use ",
            ("Export Excel", True),
            ", ",
            ("Export PDF summary", True),
            ", or ",
            ("Export PPT (checker)", True),
            ". Files download to your browser.",
        ),
    ),
    (
        "bullet",
        R("MIS flows never write to the General Ledger or GST registers."),
    ),
    (
        "bullet",
        R("PPT export typically requires a quality score of at least 70 and a separate checker."),
    ),

    ("h2", "24.10 Advisory and data boundaries"),
    ("bullet", R("OfficeMitra reads companion modules through connectors only — it is not the books of record.")),
    ("bullet", R("Figures in MIS packs come from imported or connector-sourced facts; narrative AI must cite those facts.")),
    ("bullet", R("Without an AI provider key, task generation and briefs still work with deterministic fallbacks where implemented.")),
    ("bullet", R("Contact your platform administrator to enable office_ai, MIS, or write-back flags for your tenant.")),
]


def _set_numbering(p, num_id, ilvl=0):
    pPr = p._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    e_ilvl = OxmlElement("w:ilvl")
    e_ilvl.set(qn("w:val"), str(ilvl))
    e_num = OxmlElement("w:numId")
    e_num.set(qn("w:val"), str(num_id))
    numPr.append(e_ilvl)
    numPr.append(e_num)
    pPr.append(numPr)


def _runs(p, runs):
    if isinstance(runs, str):
        runs = [(runs, False)]
    for text, bold in runs:
        r = p.add_run(text)
        if bold:
            r.bold = True


def _heading_exists(doc: Document, text: str) -> bool:
    for para in doc.paragraphs:
        if para.style and para.style.name == "Heading 1" and para.text.strip() == text:
            return True
    return False


def build_and_insert(doc: Document) -> bool:
    if _heading_exists(doc, "24. OfficeMitra AI"):
        print("Section '24. OfficeMitra AI' already present — skipping insert.")
        return False

    anchor = None
    for para in doc.paragraphs:
        if para.style and para.style.name == "Heading 1" and para.text.strip() == "24. Workflow Tips":
            anchor = para
            break
    if anchor is None:
        raise SystemExit("Could not find the '24. Workflow Tips' heading to insert before.")

    new_paras = []
    for kind, runs in CONTENT:
        if kind == "h1":
            p = doc.add_paragraph(style="Heading 1")
            _runs(p, runs)
        elif kind == "h2":
            p = doc.add_paragraph(style="Heading 2")
            _runs(p, runs)
        elif kind == "p":
            p = doc.add_paragraph(style="Normal")
            _runs(p, runs)
        elif kind == "step":
            p = doc.add_paragraph(style="List Paragraph")
            _runs(p, runs)
            _set_numbering(p, STEP_NUMID)
        elif kind == "bullet":
            p = doc.add_paragraph(style="List Paragraph")
            _runs(p, runs)
            _set_numbering(p, BULLET_NUMID)
        else:
            continue
        new_paras.append(p)

    for p in new_paras:
        anchor._p.addprevious(p._p)

    renumber = {
        "24. Workflow Tips": "25. Workflow Tips",
        "25. Common Errors & Fixes": "26. Common Errors & Fixes",
        "26. Glossary": "27. Glossary",
        "27. Support": "28. Support",
    }
    for para in doc.paragraphs:
        if para.style and para.style.name == "Heading 1":
            t = para.text.strip()
            if t in renumber and para.runs:
                para.runs[0].text = renumber[t]
                for r in para.runs[1:]:
                    r.text = ""
    return True


def emit_markdown() -> None:
    lines = [
        "# MitraBooks — OfficeMitra AI (User Guide)",
        "",
        "*Source for section 24 added to `docs/MitraBooks_User_Manual.docx`.*",
        "",
    ]
    for kind, runs in CONTENT:
        if isinstance(runs, str):
            runs = [(runs, False)]
        text = "".join((f"**{t}**" if b else t) for t, b in runs)
        if kind == "h1":
            lines += ["", f"## {text}", ""]
        elif kind == "h2":
            lines += ["", f"### {text}", ""]
        elif kind == "p":
            lines += [text, ""]
        elif kind == "step":
            lines.append(f"1. {text}")
        elif kind == "bullet":
            lines.append(f"- {text}")
    MD_OUT.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"Wrote {MD_OUT.relative_to(ROOT)}")


def main() -> None:
    doc = Document(str(MANUAL))
    changed = build_and_insert(doc)
    if changed:
        doc.save(str(MANUAL))
        print(f"Updated {MANUAL.relative_to(ROOT)}")
    if ASSET_COPY.parent.exists():
        import shutil

        try:
            shutil.copyfile(MANUAL, ASSET_COPY)
            print(f"Synced {ASSET_COPY.relative_to(ROOT)}")
        except PermissionError:
            print(f"Note: Could not overwrite {ASSET_COPY.name} — close Word and re-run to sync.")
    emit_markdown()


if __name__ == "__main__":
    main()
